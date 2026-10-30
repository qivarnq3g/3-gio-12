import argparse
import hashlib
import json
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DESIGN_DIR = ROOT / "docs" / "thiet-ke"
MASTER = "thuyet-minh.md"
PARTS = [
    ("kien-truc-he-thong.md", 10, "Kiến trúc hệ thống"),
    ("luoc-do-use-case.md", 9, "Lược đồ use case"),
    ("luoc-do-lop-game.md", 11, "Lược đồ lớp phần game"),
    ("luoc-do-lop-may-chu.md", 12, "Lược đồ lớp phần máy chủ và web"),
]
FONT = "Times New Roman"
BROWSERS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
]
SCALE = 2
PORTRAIT = (16.0, 22.0)
LANDSCAPE = (24.5, 14.5)
MERMAID_CONFIG = {
    "fontFamily": "Arial",
    **{kind: {"useMaxWidth": False} for kind in ("flowchart", "sequence", "class", "er", "state")},
}
PAGE_MARGINS = '<w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1701" w:header="567" w:footer="567" w:gutter="0"/>'
LANDSCAPE_START = f"""```{{=openxml}}
<w:p><w:pPr><w:sectPr><w:pgSz w:w="11906" w:h="16838"/>{PAGE_MARGINS}</w:sectPr></w:pPr></w:p>
```"""
LANDSCAPE_END = f"""```{{=openxml}}
<w:p><w:pPr><w:sectPr><w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/>{PAGE_MARGINS}</w:sectPr></w:pPr></w:p>
```"""
HEADING = re.compile(r"^(#{1,6})\s+(?:\d+(?:\.\d+)*\.?\s+)?(.*?)\s*$")
FENCE = re.compile(r"^```\s*mermaid\s*$")
TOC_XML = """```{=openxml}
<w:p><w:r><w:br w:type="page"/></w:r></w:p>
<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Mục lục</w:t></w:r></w:p>
<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Nhấn F9 để cập nhật mục lục.</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
<w:p><w:r><w:br w:type="page"/></w:r></w:p>
```"""


def split_cover(text):
    lines = text.splitlines()
    for i, line in enumerate(lines):
        if line.startswith("# "):
            return "\n".join(lines[:i]), "\n".join(lines[i:])
    return "", text


def drop_intro(text):
    lines = text.splitlines()
    out, skipping = [], False
    for line in lines:
        if line.startswith("# "):
            out.append(line)
            skipping = True
            continue
        if skipping and line.startswith("## "):
            skipping = False
        if not skipping:
            out.append(line)
    return "\n".join(out)


def number_headings(text, chapter_start):
    counters = [chapter_start - 1, 0, 0, 0, 0, 0]
    out, in_code = [], False
    for line in text.splitlines():
        if line.startswith("```"):
            in_code = not in_code
        match = None if in_code else HEADING.match(line)
        if not match:
            out.append(line)
            continue
        level = len(match.group(1))
        counters[level - 1] += 1
        for k in range(level, len(counters)):
            counters[k] = 0
        label = ".".join(str(c) for c in counters[:level])
        out.append(f"{match.group(1)} {label}. {match.group(2)}")
    return "\n".join(out), counters[0]


def rewrite_references(text, chapter, chapters):
    for name, number in chapters.items():
        text = re.sub(r"\[`?" + re.escape(name) + r"`?\]\(" + re.escape(name) + r"[^)]*\)", f"chương {number}", text)
        text = text.replace(f"`{name}`", f"chương {number}")
    for name, task, _ in PARTS:
        text = re.sub(rf"task #{task}\b", f"chương {chapters[name]}", text)
    text = re.sub(r"\s*\((?:đã có từ |đã làm ở )?task #\d+(?:, #\d+)*\)", "", text)
    text = re.sub(r"chốt ở task #\d+, ", "", text)
    text = re.sub(r", quyết ở task #\d+", "", text)
    text = re.sub(r"(chương \d+), \1\b", r"\1", text)
    text = re.sub(r"(chương \d+) \(\1\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\((?!https?://)[^)]+\)", r"\1", text)
    architecture = chapters["kien-truc-he-thong.md"]

    def section(match):
        if match.group(2):
            return f"mục {architecture}.{match.group(1)}"
        return f"mục {chapter}.{match.group(1)}"

    return re.sub(r"(?<!bản 5, )mục (\d+(?:\.\d+)*)( của tài liệu kiến trúc)?", section, text)


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    return struct.unpack(">II", head[16:24])


def figure_size(pixel_width, pixel_height):
    width = pixel_width / SCALE * 2.54 / 96
    height = pixel_height / SCALE * 2.54 / 96
    portrait = min(1, PORTRAIT[0] / width, PORTRAIT[1] / height)
    if portrait < 0.8 and width / height >= 1.5:
        landscape = min(1, LANDSCAPE[0] / width, LANDSCAPE[1] / height)
        if landscape > portrait:
            return width * landscape, True
    return width * portrait, False


def browser_config(work):
    path = os.environ.get("MERMAID_BROWSER") or next((p for p in BROWSERS if Path(p).exists()), None)
    if not path:
        return None
    config = work / "puppeteer.json"
    config.write_text(json.dumps({"executablePath": path, "headless": "shell"}), encoding="utf-8")
    return config


def mermaid_command():
    found = shutil.which("mmdc")
    if found:
        return [found]
    npx = shutil.which("npx")
    if not npx:
        sys.exit("Cần Mermaid CLI: npm install -g @mermaid-js/mermaid-cli")
    return [npx, "-y", "@mermaid-js/mermaid-cli"]


def render_diagrams(text, chapter, work, cache):
    out, block, figure, title = [], None, 0, ""
    for line in text.splitlines():
        heading = HEADING.match(line)
        if block is None and heading:
            title = re.sub(r"^\d+(?:\.\d+)*\.\s*", "", line.lstrip("#").strip())
        if block is None and FENCE.match(line):
            block = []
            continue
        if block is not None and line.startswith("```"):
            source = "\n".join(block) + "\n"
            settings = json.dumps(MERMAID_CONFIG, sort_keys=True)
            digest = hashlib.sha256((settings + source).encode("utf-8")).hexdigest()[:16]
            image = cache / f"{digest}.png"
            if not image.exists():
                src = work / f"{digest}.mmd"
                src.write_text(source, encoding="utf-8")
                config_file = work / "mermaid.json"
                config_file.write_text(settings, encoding="utf-8")
                cmd = mermaid_command() + ["-i", str(src), "-o", str(image), "-s", str(SCALE), "-b", "white", "-c", str(config_file)]
                browser = browser_config(work)
                if browser:
                    cmd += ["-p", str(browser)]
                subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL)
            figure += 1
            size, landscape = figure_size(*png_size(image))
            line = f"![Hình {chapter}.{figure}. {title}]({image.as_posix()}){{width={size:.1f}cm}}"
            out.extend([LANDSCAPE_START, "", line, "", LANDSCAPE_END] if landscape else [line])
            block = None
            continue
        if block is not None:
            block.append(line)
        else:
            out.append(line)
    return "\n".join(out)


def placeholder(title):
    return f"# {title}\n\n## Nội dung\n\nPhần này đang được soạn.\n"


def build_markdown(work, cache):
    cover, body = split_cover((DESIGN_DIR / MASTER).read_text(encoding="utf-8"))
    chapters, missing, next_chapter = {}, [], None
    body_numbered, last = number_headings(body, 1)
    next_chapter = last + 1
    for name, _, _ in PARTS:
        chapters[name] = next_chapter
        next_chapter += 1
    pieces = [cover, TOC_XML, body_numbered]
    for name, task, title in PARTS:
        path = DESIGN_DIR / name
        if path.exists():
            text = drop_intro(path.read_text(encoding="utf-8"))
        else:
            text = placeholder(title)
            missing.append(f"{name} (task #{task})")
        chapter = chapters[name]
        text, _ = number_headings(text, chapter)
        text = rewrite_references(text, chapter, chapters)
        text = render_diagrams(text, chapter, work, cache)
        pieces.append(text)
    combined = "\n\n".join(pieces) + "\n"
    leftovers = sorted(set(re.findall(r"task #\d+", combined)) - {f"task #{t}" for n, t, _ in PARTS if f"{n} (task #{t})" in missing})
    return combined, missing, leftovers


def set_font(style, size=None, bold=None, color=True):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(attr), None)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False


def make_reference(path):
    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(str(path))
    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Table", "Caption", "Image Caption"):
        if name in [s.name for s in styles]:
            set_font(styles[name], 13)
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13), ("Heading 4", 13), ("TOC Heading", 16)):
        if name in [s.name for s in styles]:
            set_font(styles[name], size, bold=True)
    for name in ("Body Text", "First Paragraph"):
        fmt = styles[name].paragraph_format
        fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(0), Pt(6), 1.3
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    for name in ("Caption", "Image Caption"):
        if name in [s.name for s in styles]:
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name, size, bold, space in (("Cover Org", 13, True, 6), ("Cover Title", 22, True, 24), ("Cover Subtitle", 14, False, 24), ("Cover Text", 13, False, 12)):
        style = styles.add_style(name, 1)
        style.base_style = styles["Normal"]
        set_font(style, size, bold=bold)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_after = Pt(space)
        style.paragraph_format.space_before = Pt(space if name == "Cover Title" else 0)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(56.7)
        section.left_margin = Pt(85)
        section.right_margin = Pt(56.7)
    doc.save(str(path))


def finish_docx(path):
    doc = Document(str(path))
    doc.paragraphs[0].style = doc.styles["Cover Org"]
    for table in doc.tables:
        rows = table.rows
        if not rows:
            continue
        trpr = rows[0]._tr.get_or_add_trPr()
        if trpr.find(qn("w:tblHeader")) is None:
            trpr.append(OxmlElement("w:tblHeader"))
        for cell in table._tbl.iter(qn("w:shd")):
            cell.getparent().remove(cell)
        for row_index, row in enumerate(rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = row_index == 0
    first = doc.sections[0]
    first.different_first_page_header_footer = True
    footer = first.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        run = OxmlElement("w:r")
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        run.append(node)
        footer._p.append(run)
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)
    doc.save(str(path))


def main():
    parser = argparse.ArgumentParser(description="Ghép các phần thiết kế thành bản thuyết minh .docx")
    parser.add_argument("-o", "--output", default=str(DESIGN_DIR / "build" / "thuyet-minh-thiet-ke.docx"))
    parser.add_argument("--markdown", help="Ghi thêm bản Markdown đã ghép ra đường dẫn này")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    cache = output.parent / "diagrams"
    cache.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        combined, missing, leftovers = build_markdown(work, cache)
        source = work / "combined.md"
        source.write_text(combined, encoding="utf-8")
        if args.markdown:
            Path(args.markdown).write_text(combined, encoding="utf-8")
        reference = work / "reference.docx"
        make_reference(reference)
        subprocess.run(["pandoc", str(source), "-f", "markdown-auto_identifiers", "-o", str(output), "--reference-doc", str(reference), "--resource-path", str(ROOT), "--columns", "10000"], check=True)
    finish_docx(output)
    print(f"Đã ghi {output}")
    for item in missing:
        print(f"Thiếu phần: {item}")
    for item in leftovers:
        print(f"Còn nhắc tới {item} trong nội dung, cần sửa ở tệp nguồn")
    return 1 if leftovers else 0


if __name__ == "__main__":
    sys.exit(main())
